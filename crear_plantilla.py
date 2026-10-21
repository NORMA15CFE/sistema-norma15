import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generar_plantilla_oficial():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Informe Norma 15"

    # Asegurar visibilidad de cuadrícula
    ws.views.sheetView[0].showGridLines = True

    # Estilos base
    font_titulo = Font(name="Arial", size=11, bold=True)
    fill_gris = PatternFill(start_color="595959", end_color="595959", fill_type="solid")
    font_blanco = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    
    border_celda = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'),
        bottom=Side(style='thin', color='A6A6A6')
    )

    # 1. TÍTULO PRINCIPAL DE LA CFE
    ws.merge_cells('A3:L3')
    ws['A3'] = "INFORME DE SERVICIO SOCIAL, PRÁCTICAS, RESIDENCIAS O ESTADÍAS PROFESIONALES"
    ws['A3'].font = font_titulo
    ws['A3'].alignment = align_center

    # 2. ENCABEZADOS DE LA TABLA (En el orden exacto que solicitaste)
    encabezados = [
        "No.", "AÑO", "Mes reportado", "Coordinación / Gerencia", 
        "Centro de trabajo", "Tipo", "Nombre de la Persona Estudiante", 
        "Nombre de la institución", "Carrera", "Fecha de Inicio", 
        "Fecha de conclusión", "Monto del Apoyo"
    ]

    ws.row_dimensions[5].height = 25
    for col_idx, text in enumerate(encabezados, start=1):
        cell = ws.cell(row=5, column=col_idx, value=text)
        cell.fill = fill_gris
        cell.font = font_blanco
        cell.alignment = align_center
        cell.border = border_celda

    # Guardar plantilla base
    wb.save("plantilla_cfe.xlsx")
    print("¡Plantilla oficial creada con éxito!")

if __name__ == '__main__':
    generar_plantilla_oficial()
    